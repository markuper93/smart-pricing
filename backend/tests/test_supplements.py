"""Run: venv/Scripts/python.exe -m unittest discover -s tests -p test_supplements.py -v
Always uses an isolated temporary SQLite database, never the configured DB.
Optional real-file test: SUPPLEMENTS_TEST_XLSX=/path/to/vendor.xlsx.
"""
import base64
import io
import json
import os
from pathlib import Path
import tempfile
import unittest

_TEMP = tempfile.TemporaryDirectory(prefix='smart-pricing-supplements-')
os.environ['DATABASE_URL'] = 'sqlite:///' + (Path(_TEMP.name) / 'test.db').as_posix()
os.environ['SECRET_KEY'] = 'isolated-supplements-tests-only'
os.environ['ADMIN_PASS'] = 'isolated-tests-only'
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, text, inspect
from app.database import Base, engine, SessionLocal
from app.models.user import User, TrackingGroup, TrackingItem
from app.models.price_data import PriceList, PriceEntry, Supplement
from app.routes import admin, user, reports
from app.utils.auth import create_access_token
from app.utils.supplements import parse_supplements, stable_key
from app.migrations import migrate_supplements

app = FastAPI()
for router in (admin.router, user.router, reports.router):
    app.include_router(router)
client = TestClient(app)


def workbook(rows):
    wb = Workbook()
    wb.active.append(['KOD', 'YEAR0', 'TEUR', 'AHUZ', 'COST', 'YEAR1', 'YEAR2', 'LETTER', 'LETTER', 'ID4WS'])
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sample(percent=5):
    return workbook([
        [929200, 2024, 'NOBLE', percent, 999999, 2022, 2024, 'a', 'x', 123],
        [929200, 2025, 'ULTIMATE', 2, 888888, 2024, 2025, 'b', 'x', 124],
        [929200, 2025, 'L/F NOBLE', 0, 888888, 2025, 2025, 'c', 'x', 125],
        [929200, 2025, 'הערה', None, None, 2022, 2025, 'd', 'x', 126],
        [929200, 2024, 'גג שמש', 1, 888888, 2022, 2025, 'e', 'x', 127],
        [929200, 2024, 'ריפודי עור', -2, 888888, 2022, 2025, 'f', 'x', 128],
    ])


class SupplementsTests(unittest.TestCase):
    def setUp(self):
        Base.metadata.drop_all(engine)
        Base.metadata.create_all(engine)
        with SessionLocal() as db:
            for id_, adm in [(1, True), (2, False), (3, False)]:
                db.add(User(id=id_, email=f'{id_}@test.invalid', username=f'u{id_}', hashed_password='unused', is_admin=adm))
            # ID order deliberately differs from chronological order.
            db.add_all([PriceList(id=1, filename='b.csv', month='נובמבר', year=2026, label='נובמבר 2026'),
                        PriceList(id=2, filename='a.csv', month='אוקטובר', year=2026, label='אוקטובר 2026')])
            for pl in (1, 2):
                for year in (2024, 2025):
                    db.add(PriceEntry(price_list_id=pl, car_code='929200', year=year, price=148000, road_tax=5000))
            db.commit()
        self.admin = self.headers(1)
        self.user = self.headers(2)
        self.upload(1, sample())
        self.upload(2, sample(3))

    def headers(self, uid):
        return {'Authorization': 'Bearer ' + create_access_token({'sub': str(uid)})}

    def upload(self, id_, content):
        response = client.post(f'/api/admin/price-lists/{id_}/supplements', headers=self.admin,
                               files={'file': ('supplements.xlsx', content)})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def item(self, selections=None):
        return {'car_code': '929200', 'years': [2024, 2025], 'supplements': selections or {}}

    def create(self, selections=None):
        response = client.post('/api/user/groups', headers=self.user, json={'name': 'בדיקה', 'items': [self.item(selections)]})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['id']

    def compare(self, gid):
        response = client.post('/api/reports/compare', headers=self.user, json={'group_id': gid, 'month_a_id': 2, 'month_b_id': 1})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['data']

    def test_catalog_auth_chronology(self):
        url = '/api/user/car-codes/929200/supplements?year=2024'
        self.assertEqual(client.get(url).status_code, 401)
        data = client.get(url, headers=self.user).json()
        self.assertEqual(data['price_list_id'], 1)
        self.assertTrue(data['available'])
        self.assertEqual(len(data['items']), 5)
        self.assertFalse(next(r for r in data['items'] if r['kind'] == 'note')['selectable'])
        self.assertEqual(client.post('/api/admin/price-lists/1/supplements', headers=self.user, files={'file': ('a.xlsx', sample())}).status_code, 403)
        self.assertEqual(client.get(url + '&price_list_id=999', headers=self.user).status_code, 404)
        data = client.get('/api/admin/price-lists', headers=self.admin).json()
        self.assertTrue(all(r['supplements_uploaded'] and r['supplements_count'] == 6 for r in data))

    def test_selection_validation_atomic(self):
        noble, ultimate, note, zero = [stable_key('929200', d) for d in ('NOBLE', 'ULTIMATE', 'הערה', 'L/F NOBLE')]
        invalid = [{'2024': [noble, ultimate]}, {'2024': [note]}, {'2024': [zero]}, {'2023': [noble]}, {'2024': ['unknown']}, {'2024': [noble, noble]}]
        for selections in invalid:
            response = client.post('/api/user/groups', headers=self.user, json={'name': 'bad', 'items': [self.item(selections)]})
            self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(client.get('/api/user/groups', headers=self.user).json(), [])
        gid = self.create({'2024': [noble]})
        response = client.put(f'/api/user/groups/{gid}', headers=self.user, json={'items': [self.item(invalid[0])]})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(client.get('/api/user/groups', headers=self.user).json()[0]['items'][0]['supplements'], {'2024': [noble]})

    def test_per_year_persistence_updates_templates_and_ownership(self):
        selections = {'2024': [stable_key('929200', 'NOBLE')], '2025': [stable_key('929200', 'L/F NOBLE')]}
        gid = self.create(selections)
        self.assertEqual(client.get('/api/user/groups', headers=self.user).json()[0]['items'][0]['supplements'], selections)
        response = client.post(f'/api/user/groups/{gid}/items', headers=self.headers(3), json=self.item())
        self.assertEqual(response.status_code, 404)
        self.assertEqual(client.put(f'/api/user/groups/{gid}', headers=self.user, json={'items': [self.item(selections)]}).status_code, 200)
        self.assertEqual(client.post(f'/api/user/groups/{gid}/items', headers=self.user, json=self.item(selections)).json()['supplements'], selections)
        self.assertEqual(client.post('/api/admin/default-template', headers=self.admin, json={'items': [self.item(selections)]}).status_code, 200)
        self.assertEqual(client.get('/api/admin/default-template', headers=self.admin).json()['items'][0]['supplements'], selections)
        self.assertEqual(client.post('/api/admin/users/3/inject-defaults', headers=self.admin).status_code, 200)
        self.assertEqual(client.get('/api/user/groups', headers=self.headers(3)).json()[0]['items'][0]['supplements'], selections)

    def test_report_percentage_changes_zero_accessories_and_no_tax(self):
        selected = {'2024': [stable_key('929200', 'NOBLE')], '2025': [stable_key('929200', 'L/F NOBLE')]}
        gid = self.create(selected)
        rows = self.compare(gid)
        self.assertEqual(rows[0]['בסיס + תוספות'], '155,400 (+2,960)')
        self.assertIn('5%', rows[0]['תוספות'])
        self.assertIn('+2 נקודות אחוז', rows[0]['תוספות'])
        self.assertEqual(rows[1]['בסיס + תוספות'], '148,000 (+0)')
        self.assertIn('0%', rows[1]['תוספות'])
        selected['2024'] += [stable_key('929200', 'גג שמש'), stable_key('929200', 'ריפודי עור')]
        client.post(f'/api/user/groups/{gid}/items', headers=self.user, json=self.item(selected))
        self.assertEqual(self.compare(gid)[0]['בסיס + תוספות'], '153,920 (+2,960)')

    def test_missing_upload_record_and_legacy(self):
        key = stable_key('929200', 'NOBLE')
        gid = self.create({'2024': [key]})
        with SessionLocal() as db:
            db.get(PriceList, 2).supplements_uploaded = False
            db.commit()
        self.assertEqual(self.compare(gid)[0]['בסיס + תוספות'], '155,400 (שינוי: לא ידוע)')
        with SessionLocal() as db:
            db.query(Supplement).filter(Supplement.price_list_id == 1, Supplement.key == key).delete()
            db.commit()
        self.assertEqual(self.compare(gid)[0]['בסיס + תוספות'], 'לא ידוע')
        with SessionLocal() as db:
            db.get(PriceList, 1).supplements_uploaded = False
            db.commit()
        self.assertFalse(client.get('/api/user/car-codes/929200/supplements?year=2024', headers=self.user).json()['available'])
        legacy = self.create()
        row = self.compare(legacy)[0]
        self.assertEqual(row['בסיס + תוספות'], '148,000')
        self.assertEqual(row['שינוי באחוז'], '0.00%')
        self.assertEqual(row['תוספות'], 'ללא תוספות')

    def test_replace_stable_key_and_invalid_upload_atomic(self):
        before = client.get('/api/user/car-codes/929200/supplements?year=2024', headers=self.user).json()['items']
        self.upload(1, sample(6))
        after = client.get('/api/user/car-codes/929200/supplements?year=2024', headers=self.user).json()['items']
        self.assertEqual({r['key'] for r in before}, {r['key'] for r in after})
        response = client.post('/api/admin/price-lists/1/supplements', headers=self.admin, files={'file': ('bad.xlsx', b'bad')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(client.get('/api/user/car-codes/929200/supplements?year=2024', headers=self.user).json()['items'], after)

    def test_existing_selections_survive_catalog_changes(self):
        selected = {'2024': [stable_key('929200', 'NOBLE')]}
        gid = self.create(selected)
        other_gid = self.create()
        template_url = '/api/admin/default-template'
        self.assertEqual(client.post(template_url, headers=self.admin, json={'items': [self.item(selected)]}).status_code, 200)
        writers = [
            lambda item: client.put(f'/api/user/groups/{gid}', headers=self.user, json={'items': [item]}),
            lambda item: client.post(f'/api/user/groups/{gid}/items', headers=self.user, json=item),
            lambda item: client.post(template_url, headers=self.admin, json={'items': [item]}),
        ]
        # Exercise both a missing key and a newest month with no upload at all.
        for missing_upload in (False, True):
            with SessionLocal() as db:
                db.query(Supplement).filter_by(price_list_id=1, key=selected['2024'][0]).delete()
                db.get(PriceList, 1).supplements_uploaded = not missing_upload
                db.commit()
            for index, write in enumerate(writers):
                with self.subTest(missing_upload=missing_upload, writer=index):
                    response = write(self.item(selected))
                    self.assertEqual(response.status_code, 200, response.text)
                    for bad in (
                        self.item({'2024': selected['2024'] + ['unknown']}),
                        self.item({'2025': selected['2024']}),
                        dict(self.item(selected), car_code='other'),
                        self.item({'2024': selected['2024'] * 2}),
                    ):
                        self.assertEqual(write(bad).status_code, 422)
            self.assertEqual(client.get('/api/user/groups', headers=self.user).json()[0]['items'][0]['supplements'], selected)
            self.assertEqual(client.get(template_url, headers=self.admin).json()['items'][0]['supplements'], selected)
            self.assertEqual(client.post('/api/user/groups', headers=self.user, json={'name': 'new', 'items': [self.item(selected)]}).status_code, 422)
            self.assertEqual(client.put(f'/api/user/groups/{other_gid}', headers=self.user, json={'items': [self.item(selected)]}).status_code, 422)
            self.assertEqual(client.post(f'/api/user/groups/{other_gid}/items', headers=self.user, json=self.item(selected)).status_code, 422)
        # Removed keys lose their preservation privilege.
        for write in writers:
            self.assertEqual(write(self.item()).status_code, 200)
            self.assertEqual(write(self.item(selected)).status_code, 422)

    def test_retained_selection_can_add_valid_new_accessory(self):
        selected = {'2024': [stable_key('929200', 'NOBLE')]}
        gid = self.create(selected)
        with SessionLocal() as db:
            db.query(Supplement).filter_by(price_list_id=1, key=selected['2024'][0]).delete()
            db.commit()
        selected['2024'].append(stable_key('929200', 'גג שמש'))
        response = client.put(f'/api/user/groups/{gid}', headers=self.user, json={'items': [self.item(selected)]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(client.get('/api/user/groups', headers=self.user).json()[0]['items'][0]['supplements'], selected)

    def test_all_conflicting_upload_preserves_previous_records(self):
        url = '/api/user/car-codes/929200/supplements?year=2024'
        before = client.get(url, headers=self.user).json()
        for bounds in ((2022, 2024), (2023, 2025)):
            with self.subTest(bounds=bounds):
                content = workbook([
                    [929200, 2024, 'NOBLE', 5, None, 2022, 2024],
                    [929200, 2024, 'NOBLE', 8, None, *bounds],
                ])
                diagnostics = {}
                self.assertEqual(parse_supplements(content, diagnostics), [])
                self.assertEqual(diagnostics, {'skipped_rows': 0, 'ambiguous_keys': 1})
                response = client.post('/api/admin/price-lists/1/supplements', headers=self.admin,
                                       files={'file': ('conflicts.xlsx', content)})
                self.assertEqual(response.status_code, 400, response.text)
                self.assertEqual(client.get(url, headers=self.user).json(), before)
                with SessionLocal() as db:
                    self.assertEqual(db.query(Supplement).filter_by(price_list_id=1).count(), 6)
                    self.assertTrue(db.get(PriceList, 1).supplements_uploaded)

    def test_exports(self):
        gid = self.create({'2024': [stable_key('929200', 'NOBLE'), stable_key('929200', 'גג שמש')]})
        payload = {'group_id': gid, 'month_a_id': 2, 'month_b_id': 1}
        excel = client.post('/api/reports/export/excel', headers=self.user, json=payload)
        self.assertEqual(excel.status_code, 200, excel.text)
        wb = load_workbook(io.BytesIO(base64.b64decode(excel.json()['data'])))
        headers = [c.value for c in wb.active[1]]
        self.assertIn('בסיס + תוספות', headers)
        cell = wb.active.cell(2, headers.index('תוספות') + 1)
        self.assertTrue(cell.alignment.wrap_text)
        self.assertIn('\n', cell.value)
        pdf = client.post('/api/reports/export/pdf', headers=self.user, json=payload)
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_parser_quarantines_and_vendor_id_independence(self):
        diagnostics = {}
        data = workbook([[1, 2024, 'trim', 5, 100, 2022, 2024, None, None, 1],
                         [1, 2024, 'trim', 8, 100, 2023, 2025, None, None, 2],
                         [2, 2024, 'bad', 3, 100, 2025, 2020, None, None, 3]])
        self.assertEqual(parse_supplements(data, diagnostics), [])
        self.assertEqual(diagnostics, {'skipped_rows': 1, 'ambiguous_keys': 1})
        a = parse_supplements(workbook([[1, 2024, ' Trim  X ', 5, 100, 2022, 2024, None, None, 1]]))[0]
        b = parse_supplements(workbook([[1, 2024, 'trim x', 6, 999, 2022, 2025, None, None, 999]]))[0]
        self.assertEqual(a['key'], b['key'])

    def test_migration_twice_preserves_existing_rows(self):
        isolated = create_engine('sqlite://')
        with isolated.begin() as c:
            c.execute(text('CREATE TABLE price_lists (id INTEGER PRIMARY KEY)'))
            c.execute(text('CREATE TABLE tracking_items (id INTEGER PRIMARY KEY)'))
            c.execute(text('INSERT INTO price_lists VALUES (1)'))
            c.execute(text('INSERT INTO tracking_items VALUES (1)'))
        migrate_supplements(isolated)
        migrate_supplements(isolated)
        with isolated.connect() as c:
            self.assertEqual(c.execute(text('SELECT supplements_uploaded FROM price_lists')).scalar(), 0)
            self.assertEqual(c.execute(text('SELECT supplements FROM tracking_items')).scalar(), '{}')
        isolated.dispose()

    @unittest.skipUnless(os.getenv('SUPPLEMENTS_TEST_XLSX'), 'set SUPPLEMENTS_TEST_XLSX for real workbook')
    def test_real_workbook_upload(self):
        summary = self.upload(1, Path(os.environ['SUPPLEMENTS_TEST_XLSX']).read_bytes())
        print('REAL_WORKBOOK_SUMMARY', summary)
        csv_path = Path(os.environ['SUPPLEMENTS_TEST_XLSX']).with_name('מחירון נובמבר 26.csv')
        if csv_path.exists():
            from app.utils.csv_parser import parse_csv
            real = next(r for r in parse_csv(str(csv_path)) if r['car_code'] == '929200' and r['year'] == 2024)
            self.assertEqual(real['price'], 148000)
            with SessionLocal() as db:
                pe = db.query(PriceEntry).filter_by(price_list_id=1, car_code='929200', year=2024).one()
                pe.price, pe.road_tax = real['price'], real['road_tax']
                db.commit()
        items = client.get('/api/user/car-codes/929200/supplements?year=2024', headers=self.user).json()['items']
        noble = next(r for r in items if r['percent'] == 5)
        self.assertEqual(len([r for r in items if r['selectable']]), 2)
        gid = self.create({'2024': [noble['key']]})
        self.assertTrue(self.compare(gid)[0]['בסיס + תוספות'].startswith('155,400'))

    def test_missing_base_and_delete_cleanup(self):
        gid = self.create({'2024': [stable_key('929200', 'NOBLE')]})
        with SessionLocal() as db:
            db.query(PriceEntry).filter_by(price_list_id=1, year=2024).delete()
            db.commit()
        self.assertEqual(self.compare(gid)[0]['בסיס + תוספות'], 'לא ידוע')
        response = client.delete('/api/admin/price-lists/1', headers=self.admin)
        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            self.assertEqual(db.query(Supplement).filter_by(price_list_id=1).count(), 0)

    def test_empty_valid_upload_and_bad_shape(self):
        summary = self.upload(1, workbook([]))
        self.assertEqual(summary['supplements_count'], 0)
        data = client.get('/api/user/car-codes/929200/supplements?year=2024', headers=self.user).json()
        self.assertTrue(data['available'])
        self.assertEqual(data['items'], [])
        response = client.post('/api/user/groups', headers=self.user, json={'name': 'bad', 'items': [{'years': [2024]}]})
        self.assertEqual(response.status_code, 422)


def tearDownModule():
    client.close()
    engine.dispose()
    _TEMP.cleanup()
